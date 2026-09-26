import io
import json
import subprocess
import sys
import unittest
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import gemini_rest
from cv_document import (
    extract_cv_content,
    find_github_profile_url,
    find_linkedin_profile_url,
    infer_cv_sections_from_text,
)
from docx import Document
from linkedin_evidence import normalize_profile_url


class CvDocumentTests(unittest.TestCase):
    def test_txt_preserves_unicode_and_normalizes_spaced_github_url(self):
        content = "ML Engineer ‑ NLP\nGitHub: github.com / thanujan-ml".encode("utf-8-sig")
        text, links = extract_cv_content(content, ".txt")
        self.assertIn("‑", text)
        self.assertEqual([], links)
        self.assertEqual("https://github.com/thanujan-ml", find_github_profile_url(text, links))

    def test_docx_embedded_hyperlink_is_included(self):
        document = Document()
        document.add_paragraph("GitHub profile")
        document.part.rels.add_relationship(
            "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
            "https://github.com/thanujan-ml",
            "rIdGithub",
            is_external=True,
        )
        stream = io.BytesIO()
        document.save(stream)
        text, links = extract_cv_content(stream.getvalue(), ".docx")
        self.assertIn("https://github.com/thanujan-ml", text)
        self.assertEqual("https://github.com/thanujan-ml", find_github_profile_url(text, links))

    def test_docx_tables_headers_and_footers_are_extracted(self):
        document = Document()
        document.sections[0].header.paragraphs[0].text = "Ishara Fernando"
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Education"
        table.cell(0, 1).text = "BSc Business Management"
        document.sections[0].footer.paragraphs[0].text = "Colombo, Sri Lanka"
        stream = io.BytesIO()
        document.save(stream)

        text, _ = extract_cv_content(stream.getvalue(), ".docx")

        self.assertIn("Ishara Fernando", text)
        self.assertIn("Education | BSc Business Management", text)
        self.assertIn("Colombo, Sri Lanka", text)

    def test_image_cv_is_deferred_to_visual_extraction(self):
        text, links = extract_cv_content(b"image bytes", ".png")
        self.assertEqual("", text)
        self.assertEqual([], links)

    def test_fallback_parser_extracts_languages_and_projects_from_plain_text(self):
        text = """
        Full Name: Nimal Perera
        Languages: English (Fluent), Tamil (Moderate), Sinhala (Basic)
        Projects:
        1. Smart Inventory System — Built a React + Node.js dashboard for stock management.
        2. Campus Portal — PHP, MySQL, JavaScript project for student services.
        """

        parsed = infer_cv_sections_from_text(text)

        self.assertEqual(["English", "Tamil", "Sinhala"], [item["language"] for item in parsed["languages"]])
        self.assertEqual(["Smart Inventory System", "Campus Portal"], [item["name"] for item in parsed["projects"]])

    def test_linkedin_profile_is_found_and_canonicalized(self):
        text = "LinkedIn: www.linkedin.com / in / tharmapalan-thanujan"
        self.assertEqual(
            "https://www.linkedin.com/in/tharmapalan-thanujan",
            find_linkedin_profile_url(text),
        )
        self.assertEqual(
            "https://www.linkedin.com/in/tharmapalan-thanujan",
            normalize_profile_url("linkedin.com/in/tharmapalan-thanujan/"),
        )
        self.assertIsNone(normalize_profile_url("https://www.linkedin.com/posts/example"))


class GeminiEncodingTests(unittest.TestCase):
    @patch("gemini_rest.tempfile.NamedTemporaryFile")
    @patch("gemini_rest.shutil.which", return_value="curl.exe")
    @patch("gemini_rest.subprocess.run")
    def test_curl_uses_utf8_for_unicode_prompt(self, run, _which, temp_file):
        temp_file.return_value.__enter__.return_value.name = "request.conf"
        run.return_value = subprocess.CompletedProcess([], 0, '{"candidates": []}\n200', "")
        gemini_rest._curl_post("https://example.test", "key", '{"text":"non‑breaking"}', 10)
        temp_file.return_value.__enter__.return_value.write.assert_called_with('{"text":"non‑breaking"}'.encode("utf-8"))

    @patch.object(gemini_rest, "_USE_CURL", True)
    @patch("gemini_rest._curl_post")
    def test_inline_pdf_is_base64_encoded_in_request(self, curl_post):
        curl_post.return_value = {
            "candidates": [{"content": {"parts": [{"text": "{}"}]}}]
        }

        result = gemini_rest.generate_json(
            "key", "model", "Extract CV", media_data=b"%PDF-test", media_mime_type="application/pdf"
        )

        payload = json.loads(curl_post.call_args.args[2])
        inline = payload["contents"][0]["parts"][0]["inlineData"]
        self.assertEqual("application/pdf", inline["mimeType"])
        self.assertEqual("JVBERi10ZXN0", inline["data"])
        self.assertEqual({}, result)


if __name__ == "__main__":
    unittest.main()
